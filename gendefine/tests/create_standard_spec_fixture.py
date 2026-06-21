"""Create a minimal Standard-Spec.xlsx test fixture for Phase 3 testing."""

import os
from openpyxl import Workbook

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")


def create_fixture():
    wb = Workbook()

    # Standard worksheet (attribute_value format)
    ws = wb.active
    ws.title = "Standard"
    ws.append(["StudyName", "SDTM-TEST"])
    ws.append(["StudyDescription", "SDTM Test Study"])
    ws.append(["ProtocolName", "SDTM-TEST"])
    ws.append(["StandardName", "SDTMIG"])
    ws.append(["StandardVersion", "3.4"])

    # Datasets worksheet
    ws = wb.create_sheet("Datasets")
    ws.append(["Dataset", "Description", "Class", "Subclass", "Structure",
               "Repeating", "Reference Data", "Key Variables"])
    ws.append(["DM", "Demographics", "SPECIAL PURPOSE", None,
               "One record per subject", "No", "No", "STUDYID, USUBJID"])
    ws.append(["AE", "Adverse Events", "EVENTS", "Adverse Event",
               "One record per adverse event", "Yes", "No", "STUDYID, USUBJID, AESEQ"])
    ws.append(["VS", "Vital Signs", "FINDINGS", "Vital Signs",
               "One record per vital sign per visit", "Yes", "No", "STUDYID, USUBJID, VSTESTCD, VISITNUM"])

    # Variables worksheet
    ws = wb.create_sheet("Variables")
    ws.append(["Order", "Dataset", "Variable", "Label", "Data Type", "Length",
               "Core", "Codelist", "Origin Type", "Origin Source", "Role",
               "Developer Notes", "Variant"])
    # DM variables
    ws.append([1, "DM", "STUDYID", "Study Identifier", "text", 7,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([2, "DM", "DOMAIN", "Domain Abbreviation", "text", 2,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([3, "DM", "USUBJID", "Unique Subject Identifier", "text", 20,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([4, "DM", "SEX", "Sex", "text", 1,
               "Req", "SEX", "Collected", "Subject", None, None, None])
    ws.append([5, "DM", "AGE", "Age", "integer", 3,
               "Exp", None, "Derived", None, None, "Derived from BRTHDTC and RFSTDTC", None])
    # AE variables
    ws.append([1, "AE", "STUDYID", "Study Identifier", "text", 7,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([2, "AE", "USUBJID", "Unique Subject Identifier", "text", 20,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([3, "AE", "AESEQ", "Sequence Number", "integer", 3,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([4, "AE", "AETERM", "Reported Term", "text", 200,
               "Req", None, "Collected", "Subject", None, None, None])
    ws.append([5, "AE", "AESEV", "Severity", "text", 20,
               "Exp", "AESEV", "Collected", "Subject", None, None, None])
    # VS variables
    ws.append([1, "VS", "STUDYID", "Study Identifier", "text", 7,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([2, "VS", "USUBJID", "Unique Subject Identifier", "text", 20,
               "Req", None, "Assigned", "Sponsor", None, None, None])
    ws.append([3, "VS", "VSTESTCD", "Vital Signs Test Short Name", "text", 8,
               "Req", "VSTESTCD", "Assigned", "Sponsor", None, None, None])
    ws.append([4, "VS", "VSTEST", "Vital Signs Test Name", "text", 40,
               "Req", "VSTEST", "Assigned", "Sponsor", None, None, None])
    ws.append([5, "VS", "VSSTRESN", "Numeric Result", "float", 8,
               "Exp", None, "Derived", None, None, None, None])
    ws.append([6, "VS", "VISITNUM", "Visit Number", "float", 3,
               "Exp", None, "Assigned", "Sponsor", None, None, None])

    # ValueLevel worksheet
    ws = wb.create_sheet("ValueLevel")
    ws.append(["Order", "Dataset", "Variable", "Where Clause", "Data Type",
               "Length", "Core", "Codelist", "Origin Type", "Origin Source",
               "Method"])
    ws.append([1, "VS", "VSSTRESN", "VSTESTCD EQ SYSBP", "float",
               8, "Exp", None, "Derived", None, "DERIVE_SYSBP"])
    ws.append([2, "VS", "VSSTRESN", "VSTESTCD EQ DIABP", "float",
               8, "Exp", None, "Derived", None, "DERIVE_DIABP"])
    ws.append([3, "VS", "VSSTRESN", "VSTESTCD EQ PULSE AND VISITNUM EQ 1", "float",
               8, "Exp", None, "Derived", None, None])

    # Methods worksheet
    ws = wb.create_sheet("Methods")
    ws.append(["Name", "Type", "Description", "Expression Context", "Expression Code"])
    ws.append(["DERIVE_SYSBP", "Computation", "Derived systolic blood pressure",
               "Python", "derive_sysbp(raw_value)"])
    ws.append(["DERIVE_DIABP", "Computation", "Derived diastolic blood pressure",
               None, None])

    # Comments worksheet
    ws = wb.create_sheet("Comments")
    ws.append(["Description", "Document", "Pages"])
    ws.append(["See annotated CRF for details", "LF.acrf", "1-3"])

    # Documents worksheet
    ws = wb.create_sheet("Documents")
    ws.append(["ID", "Title", "Href"])
    ws.append(["LF.acrf", "Annotated CRF", "acrf.pdf"])
    ws.append(["LF.ReviewersGuide", "Reviewers Guide", "reviewers-guide.pdf"])

    output = os.path.join(FIXTURES_DIR, "standard-spec-test.xlsx")
    wb.save(output)
    print(f"Created fixture: {output}")


if __name__ == "__main__":
    create_fixture()
