"""Generic spreadsheet reader driven by JSON mapping configuration."""

import logging

from openpyxl import load_workbook

from config_loader import MappingConfig

logger = logging.getLogger(__name__)


class SpreadsheetReader:
    """Generic spreadsheet reader driven by JSON mapping configuration."""

    def __init__(self, config: MappingConfig):
        self._config = config

    def read_workbook(self, excel_path: str) -> dict:
        """Read all configured worksheets from the workbook.

        Returns dict keyed by worksheet name, where each value is:
        - For "attribute_value" format: dict of key-value pairs
        - For "tabular" format: list of dicts (one per row, keyed by column name)
        """
        wb = load_workbook(filename=excel_path, read_only=True, data_only=True)
        result = {}
        for sheet in wb.worksheets:
            ws_config = self._config.get_worksheet_config(sheet.title)
            if ws_config is None:
                logger.warning("Skipping unmapped worksheet: %s", sheet.title)
                continue
            if ws_config["format"] == "attribute_value":
                result[sheet.title] = self._read_attribute_value_sheet(sheet, ws_config)
            else:
                result[sheet.title] = self._read_tabular_sheet(sheet, ws_config)
        wb.close()
        return result

    def _read_attribute_value_sheet(self, sheet, ws_config: dict) -> dict:
        """Read a key-value pair worksheet (e.g., Study)."""
        data = {}
        for row in sheet.iter_rows(min_row=1, min_col=1, max_col=2, values_only=True):
            if row[0] is not None:
                data[row[0]] = row[1]
        return data

    def _read_tabular_sheet(self, sheet, ws_config: dict) -> list[dict]:
        """Read a tabular worksheet with header row + data rows."""
        num_cols = sheet.max_column
        header = []
        for row in sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=num_cols, values_only=True):
            header = list(row)
        rows = []
        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=num_cols, values_only=True):
            row_dict = {}
            for col_name, value in zip(header, row):
                row_dict[col_name] = value
            rows.append(row_dict)
        return rows
